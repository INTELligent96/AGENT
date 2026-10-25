import os
import zipfile
import textwrap

PROJECT_NAME = "agent_core_orchestrator"

FILES = {
    # Исправляем предупреждение PyMuPDF
    "documents/extractors.py": """\
import os
from pathlib import Path
import pymupdf as fitz  # ИСПРАВЛЕНО: убрали предупреждение
from docx import Document
import openpyxl
import pytesseract
from PIL import Image
from .legacy import convert_legacy_ms_office

def extract_text(file_path: str, tools_config: dict) -> str:
    ext = Path(file_path).suffix.lower()

    if ext in ['.doc', '.xls']:
        if tools_config.get("use_ms_office"):
            file_path = convert_legacy_ms_office(file_path)
            ext = Path(file_path).suffix.lower()
        else:
            return f"[Ошибка] Для файлов {ext} требуется MS Office."

    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif ext == '.docx':
        return extract_docx(file_path)
    elif ext == '.xlsx':
        return extract_excel(file_path)
    elif ext == '.pdf':
        return extract_pdf(file_path, tools_config)
    elif ext in ['.png', '.jpg', '.jpeg', '.bmp', '.tiff']:
        return extract_image(file_path, tools_config)
    else:
        return f"[Ошибка] Неподдерживаемый формат: {ext}"

def extract_docx(file_path: str) -> str:
    doc = Document(file_path)
    text = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text.append(" | ".join(cell.text for cell in row.cells))
    return "\\n".join(text)

def extract_excel(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    for sheet in wb.sheetnames:
        text.append(f"--- Лист: {sheet} ---")
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_str.strip():
                text.append(row_str)
    return "\\n".join(text)

def extract_pdf(file_path: str, tools_config: dict) -> str:
    doc = fitz.open(file_path)
    text = []
    for i, page in enumerate(doc):
        page_text = page.get_text()
        if page_text.strip():
            text.append(page_text)
        else:
            pix = page.get_pixmap()
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text.append(ocr_image(img, tools_config))
    return "\\n".join(text)

def extract_image(file_path: str, tools_config: dict) -> str:
    img = Image.open(file_path)
    return ocr_image(img, tools_config)

def ocr_image(img: Image.Image, tools_config: dict) -> str:
    tess_path = tools_config.get("tesseract_path")
    if tess_path and os.path.exists(tess_path):
        pytesseract.pytesseract.tesseract_cmd = tess_path
    try:
        return pytesseract.image_to_string(img, lang='rus+eng')
    except Exception as e:
        return f"[Ошибка OCR] Проверь Tesseract: {e}"
""",

    "core/__init__.py": "",

    # Промпты для агента (на русском)
    "core/prompts.py": """\
SUMMARIZE_PROMPT = \"\"\"
Ты — эксперт по анализу документов. Сделай краткое и емкое резюме предоставленного текста на русском языке.
Выдели главные мысли, ключевые факты и выводы. Избегай "воды".
Текст документа:
{text}
\"\"\"

STRUCTURE_PROMPT = \"\"\"
Ты — аналитик данных. Извлеки из текста структурированную информацию.
Если в тексте есть таблицы, списки или перечисления, оформи их в виде Markdown-таблицы или маркированного списка.
Сохрани все числа и даты.
Текст:
{text}
\"\"\"

SPELLCHECK_PROMPT = \"\"\"
Ты — корректор русского языка. Проверь текст на орфографические, пунктуационные и грамматические ошибки.
Верни ТОЛЬКО исправленный текст, без своих комментариев и пояснений.
Текст:
{text}
\"\"\"

CODE_PROMPT = \"\"\"
Ты — опытный Python-разработчик. Напиши код для следующей задачи.
Верни только чистый код без пояснений Markdown, если не просят иное.
Задача: {task}
\"\"\"
""",

    # Сам оркестратор (связывает всё вместе)
    "core/agent.py": """\
import os
import yaml
from llm.registry import LLMRegistry
from documents.extractors import extract_text
from documents.writers import create_simple_docx, render_template
from core.prompts import SUMMARIZE_PROMPT, STRUCTURE_PROMPT, SPELLCHECK_PROMPT, CODE_PROMPT

class Agent:
    def __init__(self, config_path="config.yaml"):
        with open(config_path, "r", encoding="utf-8") as f:
            self.config = yaml.safe_load(f)

        self.tools = self.config.get("tools", {})
        self.registry = LLMRegistry(config_path)

    def summarize_file(self, file_path: str) -> str:
        \"\"\"Читает файл и делает резюме.\"\"\"
        print(f"Чтение файла: {file_path}...")
        text = extract_text(file_path, self.tools)
        if text.startswith("[Ошибка]"):
            return text

        print("Генерация резюме (local_fast)...")
        llm = self.registry.get_model_for_task("default_model")
        prompt = SUMMARIZE_PROMPT.format(text=text[:8000]) # Ограничиваем контекст для 7B
        return llm.generate(prompt)

    def check_spelling_file(self, file_path: str) -> str:
        \"\"\"Читает файл и исправляет ошибки.\"\"\"
        print(f"Чтение файла: {file_path}...")
        text = extract_text(file_path, self.tools)
        if text.startswith("[Ошибка]"):
            return text

        print("Проверка орфографии (local_fast)...")
        llm = self.registry.get_model_for_task("default_model")
        prompt = SPELLCHECK_PROMPT.format(text=text[:8000])
        return llm.generate(prompt)

    def structure_file(self, file_path: str) -> str:
        \"\"\"Структурирует данные из файла.\"\"\"
        print(f"Чтение файла: {file_path}...")
        text = extract_text(file_path, self.tools)
        if text.startswith("[Ошибка]"):
            return text

        print("Структурирование данных (local_logic)...")
        llm = self.registry.get_model_for_task("complex_reasoning")
        prompt = STRUCTURE_PROMPT.format(text=text[:8000])
        return llm.generate(prompt)

    def generate_code(self, task: str) -> str:
        \"\"\"Пишет код на Python.\"\"\"
        print("Генерация кода (local_coder)...")
        llm = self.registry.get_model_for_task("code_generation")
        prompt = CODE_PROMPT.format(task=task)
        return llm.generate(prompt)

    def save_text_to_word(self, text: str, output_path: str):
        \"\"\"Сохраняет текст в простой Word-документ.\"\"\"
        create_simple_docx(text, output_path)
        return output_path
""",

    "test_agent.py": """\
import os
from core.agent import Agent

def main():
    agent = Agent("config.yaml")

    # 1. Создадим тестовый файл (ИСПРАВЛЕНО: используем \\n вместо тройных кавычек)
    test_file = "test_article.txt"
    with open(test_file, "w", encoding="utf-8") as f:
        f.write("Вчера в Москве прошла международная конференция по искусственному интеллекту.\\n"
                "На ней выступили ведущие разработчики из России и Китая.\\n"
                "Главной темой стало внедрение локальных LLM в корпоративный сектор.\\n"
                "Ученые отметили, что количество параметров моделей растет, но требования к железу тоже.\\n"
                "По итогам конференции было решено создать открытый репозиторий с лучшими практиками.\\n")

    print("=== ТЕСТ 1: Резюме текста (fast) ===")
    summary = agent.summarize_file(test_file)
    print("Результат:\\n", summary)

    print("\\n=== ТЕСТ 2: Структурирование (logic) ===")
    struct = agent.structure_file(test_file)
    print("Результат:\\n", struct)

    print("\\n=== ТЕСТ 3: Проверка орфографии (fast) ===")
    bad_text_file = "test_bad.txt"
    with open(bad_text_file, "w", encoding="utf-8") as f:
        f.write("Привет, как дила? Вчира была харашая погода, но мы ни пошли гулять.")
    fixed = agent.check_spelling_file(bad_text_file)
    print("Исправленный текст:\\n", fixed)

    print("\\n=== ТЕСТ 4: Генерация кода (coder) ===")
    code = agent.generate_code("Напиши функцию на Python, которая считает факториал числа рекурсивно.")
    print("Код:\\n", code)

if __name__ == "__main__":
    main()
"""
}


def create_project():
    print(f"Создание проекта {PROJECT_NAME}...")
    for filepath, content in FILES.items():
        clean_content = textwrap.dedent(content)
        os.makedirs(os.path.dirname(filepath) or '.', exist_ok=True)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(clean_content)
        print(f"  Создан/обновлен: {filepath}")

    zip_name = f"{PROJECT_NAME}.zip"
    with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for filepath in FILES.keys():
            zipf.write(filepath)
    print(f"\\nГотово! Архив {zip_name} создан. Распакуй его в папку проекта.")


if __name__ == "__main__":
    create_project()