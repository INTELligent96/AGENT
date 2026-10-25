import os
from pathlib import Path
import pymupdf as fitz  # ИСПРАВЛЕНО: убрали предупреждение
from docx import Document
import openpyxl
import pytesseract
from PIL import Image
from .legacy import convert_legacy_ms_office

def extract_text(file_path: str, tools_config: dict) -> str:
    ext = Path(file_path).suffix.lower()

    if ext in ['.doc', '.xls']:
        if tools_config.get("use_ms_office"):
            file_path = convert_legacy_ms_office(file_path)
            ext = Path(file_path).suffix.lower()
        else:
            return f"[Ошибка] Для файлов {ext} требуется MS Office."

    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif ext == '.docx':
        return extract_docx(file_path)
    elif ext == '.xlsx':
        return extract_excel(file_path)
    elif ext == '.pdf':
        return extract_pdf(file_path, tools_config)
    elif ext in ['.png', '.jpg', '.jpeg', '.bmp', '.tiff']:
        return extract_image(file_path, tools_config)
    else:
        return f"[Ошибка] Неподдерживаемый формат: {ext}"

def extract_docx(file_path: str) -> str:
    doc = Document(file_path)
    text = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(text)

def extract_excel(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    for sheet in wb.sheetnames:
        text.append(f"--- Лист: {sheet} ---")
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_str.strip():
                text.append(row_str)
    return "\n".join(text)

def extract_pdf(file_path: str, tools_config: dict) -> str:
    doc = fitz.open(file_path)
    text = []
    for i, page in enumerate(doc):
        page_text = page.get_text()
        if page_text.strip():
            text.append(page_text)
        else:
            pix = page.get_pixmap()
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text.append(ocr_image(img, tools_config))
    return "\n".join(text)

def extract_image(file_path: str, tools_config: dict) -> str:
    img = Image.open(file_path)
    return ocr_image(img, tools_config)

def ocr_image(img: Image.Image, tools_config: dict) -> str:
    tess_path = tools_config.get("tesseract_path")
    if tess_path and os.path.exists(tess_path):
        pytesseract.pytesseract.tesseract_cmd = tess_path
    try:
        return pytesseract.image_to_string(img, lang='rus+eng')
    except Exception as e:
        return f"[Ошибка OCR] Проверь Tesseract: {e}"
