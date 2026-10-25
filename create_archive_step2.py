import os
import zipfile
import textwrap

PROJECT_NAME = "agent_documents_foundation"

FILES = {
    # Обновленный конфиг (добавили пути к Tesseract и LibreOffice)
    "config.yaml": """\
llm_profiles:
  local_coder:
    type: "openai_compat"
    base_url: "http://localhost:1235/v1"
    api_key: "lm-studio"
    model_name: "qwen-qwen2.5-coder-14b-instruct"
  local_logic:
    type: "openai_compat"
    base_url: "http://localhost:1235/v1"
    api_key: "lm-studio"
    model_name: "deepseek-r1-0528-qwen3-8b-abliterated"
  local_fast:
    type: "openai_compat"
    base_url: "http://localhost:1235/v1"
    api_key: "lm-studio"
    model_name: "qwen2.5-7b-instruct-uncensored"

routing:
  default_model: "local_fast"
  code_generation: "local_coder"
  complex_reasoning: "local_logic"

# ПУТИ К ВНЕШНИМ ПРОГРАММАМ (ПРОВЕРЬ ИХ!)
tools:
  tesseract_path: "C:\\\\Program Files\\\\Tesseract-OCR\\\\tesseract.exe"
  libreoffice_path: "C:\\\\Program Files\\\\LibreOffice\\\\program\\\\soffice.exe"
""",

    "documents/__init__.py": "",

    "documents/legacy.py": """\
import subprocess
import os
from pathlib import Path

def convert_to_modern(file_path: str, libreoffice_path: str) -> str:
    \"\"\"
    Конвертирует .doc -> .docx и .xls -> .xlsx через LibreOffice.
    Возвращает путь к новому файлу.
    \"\"\"
    file_path = Path(file_path)
    out_dir = file_path.parent

    # Определяем целевой формат
    if file_path.suffix.lower() == '.doc':
        target_format = 'docx'
    elif file_path.suffix.lower() == '.xls':
        target_format = 'xlsx'
    else:
        return str(file_path)

    print(f"Конвертация {file_path.name} в {target_format}...")

    # Запуск LibreOffice в headless-режиме
    try:
        subprocess.run([
            libreoffice_path,
            '--headless',
            '--convert-to', target_format,
            '--outdir', str(out_dir),
            str(file_path)
        ], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

        new_file = out_dir / f"{file_path.stem}.{target_format}"
        if new_file.exists():
            return str(new_file)
        else:
            raise FileNotFoundError("Конвертация не удалась, файл не создан.")
    except Exception as e:
        print(f"Ошибка конвертации (проверь путь к LibreOffice): {e}")
        return str(file_path)
""",

    "documents/extractors.py": """\
import os
from pathlib import Path
import fitz  # PyMuPDF
from docx import Document
import openpyxl
import pytesseract
from PIL import Image
from .legacy import convert_to_modern

def extract_text(file_path: str, tools_config: dict) -> str:
    \"\"\"Главный роутер для извлечения текста из любого файла.\"\"\"
    ext = Path(file_path).suffix.lower()

    # Легаси форматы
    if ext in ['.doc', '.xls']:
        lo_path = tools_config.get("libreoffice_path")
        if lo_path and os.path.exists(lo_path):
            file_path = convert_to_modern(file_path, lo_path)
            ext = Path(file_path).suffix.lower()
        else:
            return f"[Ошибка] Для файлов {ext} нужен LibreOffice по пути {lo_path}"

    # Роутинг по форматам
    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif ext == '.docx':
        return extract_docx(file_path)
    elif ext == '.xlsx':
        return extract_excel(file_path)
    elif ext == '.pdf':
        return extract_pdf(file_path, tools_config)
    elif ext in ['.png', '.jpg', '.jpeg', '.bmp', '.tiff']:
        return extract_image(file_path, tools_config)
    else:
        return f"[Ошибка] Неподдерживаемый формат: {ext}"

def extract_docx(file_path: str) -> str:
    doc = Document(file_path)
    text = []
    for para in doc.paragraphs:
        text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            text.append(" | ".join(cell.text for cell in row.cells))
    return "\\n".join(text)

def extract_excel(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    for sheet in wb.sheetnames:
        text.append(f"--- Лист: {sheet} ---")
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_str.strip():
                text.append(row_str)
    return "\\n".join(text)

def extract_pdf(file_path: str, tools_config: dict) -> str:
    doc = fitz.open(file_path)
    text = []
    for i, page in enumerate(doc):
        page_text = page.get_text()
        if page_text.strip():
            text.append(page_text)
        else:
            # Если текста нет (скан), рендерим в картинку и делаем OCR
            pix = page.get_pixmap()
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text.append(ocr_image(img, tools_config))
    return "\\n".join(text)

def extract_image(file_path: str, tools_config: dict) -> str:
    img = Image.open(file_path)
    return ocr_image(img, tools_config)

def ocr_image(img: Image.Image, tools_config: dict) -> str:
    tess_path = tools_config.get("tesseract_path")
    if tess_path and os.path.exists(tess_path):
        pytesseract.pytesseract.tesseract_cmd = tess_path
    try:
        return pytesseract.image_to_string(img, lang='rus+eng')
    except Exception as e:
        return f"[Ошибка OCR] Проверь установку Tesseract и языков rus/eng: {e}"
""",

    "documents/writers.py": """\
from docx import Document
from docxtpl import DocxTemplate
import os

def create_simple_docx(text: str, output_path: str):
    \"\"\"Создает простой Word-документ из текста.\"\"\"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(output_path)

def render_template(template_path: str, context: dict, output_path: str):
    \"\"\"
    Заполняет твой шаблон .docx (с сохранением стилей).
    В шаблоне должны быть метки вроде {{ title }} или {{ text }}.
    \"\"\"
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон {template_path} не найден!")

    tpl = DocxTemplate(template_path)
    tpl.render(context)
    tpl.save(output_path)
    return output_path
""",

    "test_documents.py": """\
import yaml
import os
from documents.extractors import extract_text
from documents.writers import create_simple_docx, render_template

def main():
    # Загружаем конфиг для путей к Tesseract и LibreOffice
    with open("config.yaml", "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)
    tools = config.get("tools", {})

    print("--- Тест 1: Создание и чтение Word (.docx) ---")
    test_file = "test_output.docx"
    create_simple_docx("Привет! Это тестовый документ для проверки стилей.", test_file)
    print("Создан файл:", test_file)
    text = extract_text(test_file, tools)
    print("Прочитан текст:", text)

    print("\\n--- Тест 2: Чтение Excel (.xlsx) ---")
    # Создадим тестовый Excel
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Имя", "Возраст"])
    ws.append(["Иван", 30])
    wb.save("test_excel.xlsx")
    print(extract_text("test_excel.xlsx", tools))

    print("\\n--- Тест 3: Проверка OCR (Tesseract) ---")
    print(f"Путь к Tesseract: {tools.get('tesseract_path')}")
    if os.path.exists(tools.get('tesseract_path', '')):
        print("Tesseract найден. (Тест картинки пропущен для экономии времени)")
    else:
        print("[Внимание] Tesseract не найден по указанному пути. OCR не будет работать.")

if __name__ == "__main__":
    main()
"""
}


def create_project():
    print(f"Создание проекта {PROJECT_NAME}...")
    for filepath, content in FILES.items():
        clean_content = textwrap.dedent(content)
        os.makedirs(os.path.dirname(filepath) or '.', exist_ok=True)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(clean_content)
        print(f"  Создан: {filepath}")

    zip_name = f"{PROJECT_NAME}.zip"
    with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for filepath in FILES.keys():
            zipf.write(filepath)
    print(f"\\nГотово! Архив {zip_name} создан.")


if __name__ == "__main__":
    create_project()